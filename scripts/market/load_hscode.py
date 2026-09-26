#!/usr/bin/env python3
"""관세청 HS부호 엑셀(scripts/market/관세청_HS부호_20260101.xlsx)을 읽어
instance/hscode.db의 hs0code_master 테이블에 적재한다.

sabuzak.db(instance/sabuzak.db)는 박람회(raw_exhibitions) 전용으로만 쓰기로
했으므로, 이 스크립트는 sabuzak.db를 전혀 건드리지 않고 hscode.db에 바로
적재한다 (예전 hscode_recommend/load_excel.py, make_db.py는 sabuzak.db에
넣은 뒤 scripts/split_hscode_db.py로 옮기는 2단계 방식이었는데, 그 중간
단계가 실제로는 한 번도 끝까지 실행되지 않아 hs0code_master가 계속 비어있는
원인이 됐다).

[엑셀 구조에서 발견한 문제와 해결]
관세청 엑셀은 10자리 코드(실제 신고에 쓰는 리프 코드)와 8자리 코드(그
윗단계 소분류, 한글품목명 컬럼에만 실제 품목 이름이 있고 하위 10자리들은
"건조한 것"/"냉동한 것"/"기타" 같은 속성만 적혀 있음)가 섞여 있다.
예) 8자리 "12122110" -> "김", 10자리 "1212211010" -> "건조한 것"
이 상태로 10자리 행의 한글품목명만 검색하면 "김"으로 검색해도 안 걸린다.
그래서 10자리(리프) 행마다 바로 위 8자리 행의 이름을 찾아
"김 건조한 것"처럼 합쳐서 name_ko로 저장한다.

또한 이 엑셀은 챕터 01~09(코드가 0으로 시작)인 행 일부가 숫자 셀로 저장돼
있어 앞자리 0이 사라진 채로 읽힌다(7자리/9자리로 나타남). 8/10자리로
0을 채워 정규화한다.

실행 (sabuzak 루트에서):
    pip install openpyxl   # requirements.txt에 이미 있음
    python scripts/market/load_hscode.py
    python scripts/market/load_hscode.py --excel "다른파일.xlsx" --db instance/hscode.db
"""

import argparse
import os
import sqlite3

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_EXCEL = os.path.join(BASE_DIR, "관세청_HS부호_20260101.xlsx")
DEFAULT_DB = os.path.join(BASE_DIR, "..", "..", "instance", "hscode.db")

COL_HSCODE = 0
COL_NAME_KO = 3


def _normalize_code(raw):
    code = str(raw).strip()
    if len(code) == 7:
        return code.zfill(8)
    if len(code) == 9:
        return code.zfill(10)
    return code


def load_rows(excel_path):
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_by_code8 = {}
    leaf_rows = []  # (code10, name_ko)

    rows = ws.iter_rows(values_only=True)
    next(rows, None)  # 헤더 행 스킵
    for row in rows:
        raw_code = row[COL_HSCODE]
        name_ko = row[COL_NAME_KO]
        if raw_code is None:
            continue
        code = _normalize_code(raw_code)
        name_ko = (name_ko or "").strip()

        if len(code) == 8:
            if name_ko:
                header_by_code8[code] = name_ko
        elif len(code) == 10:
            leaf_rows.append((code, name_ko))
        # 그 외 길이(정상적으로 8/10이 안 되는 손상된 행)는 건너뜀

    wb.close()

    entries = []
    for code, leaf_name in leaf_rows:
        ancestor = header_by_code8.get(code[:8])
        if ancestor and ancestor != leaf_name:
            display_name = f"{ancestor} {leaf_name}".strip()
        else:
            display_name = leaf_name or ancestor or ""
        entries.append((code, leaf_name, display_name))

    return entries


def save_entries(db_path, entries):
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.execute("DROP TABLE IF EXISTS hs0code_master")
    conn.execute(
        "CREATE TABLE hs0code_master (hscode TEXT, hsk_name TEXT, name_ko TEXT)"
    )
    conn.executemany(
        "INSERT INTO hs0code_master (hscode, hsk_name, name_ko) VALUES (?, ?, ?)",
        entries,
    )
    conn.execute("CREATE INDEX idx_hs0code_master_name_ko ON hs0code_master(name_ko)")
    conn.commit()
    conn.close()


def main():
    parser = argparse.ArgumentParser(description="관세청 HS부호 엑셀 -> instance/hscode.db 적재")
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--db", default=DEFAULT_DB)
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다: {args.excel}")

    print(f"엑셀 읽는 중: {args.excel}")
    entries = load_rows(args.excel)
    print(f"리프(10자리) 코드 {len(entries)}건 정리 완료")

    save_entries(args.db, entries)
    print(f"완료: {args.db} 의 hs0code_master에 {len(entries)}건 적재")
    print("예시:")
    for e in entries[:3]:
        print(" ", e)


if __name__ == "__main__":
    main()
